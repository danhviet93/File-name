# pip install pymupdf openpyxl
import os, sys, re, datetime, pathlib, unicodedata
import pymupdf
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

# Caractères interdits dans le XML d'Excel: contrôles, mais aussi U+FFFE/U+FFFF et "surrogates" isolés
# (texte PDF mal encodé, métadonnées). Sans lxml, openpyxl les écrit quand même -> fichier corrompu.
XML_BAD = re.compile("[^\u0009\u000A\u000D\u0020-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]")

def clean(x):
    # retire les caractères interdits dans Excel (sinon Excel "répare" le fichier et supprime la feuille)
    return XML_BAD.sub("", ILLEGAL_CHARACTERS_RE.sub("", x)) if isinstance(x, str) else x

pymupdf.TOOLS.mupdf_display_errors(False)

# ---- Dossier à analyser: argument en ligne de commande, sinon boîte de dialogue ----
# ---- Détection du cartouche et extraction du titre ----
TXT_FLAGS = (pymupdf.TEXT_PRESERVE_LIGATURES | pymupdf.TEXT_PRESERVE_WHITESPACE | pymupdf.TEXT_MEDIABOX_CLIP |
             getattr(pymupdf, "TEXT_CID_FOR_UNKNOWN_UNICODE", 0))   # comme le mode par défaut, mais SANS images (beaucoup plus rapide)
# Police sans table Unicode: code du caractère = n° de glyphe dans l'ordre standard Macintosh des polices
# TrueType (3 = espace, 36 = A, 68 = a, 112 = é, 123 = ô...). Table complète (258 glyphes) à partir du n° 3.
MAC_GLYPHS = ' !"#$%&\'()*+,-./0123456789:;<=>?@ABCDEFGHIJKLMNOPQRSTUVWXYZ[\\]^_`abcdefghijklmnopqrstuvwxyz{|}~ÄÅÇÉÑÖÜáàâäãåçéèêëíìîïñóòôöõúùûü†°¢£§•¶ß®©™´¨≠ÆØ∞±≤≥¥µ∂∑∏π∫ªºΩæø¿¡¬√ƒ≈∆«»…\xa0ÀÃÕŒœ–—“”‘’÷◊ÿŸ⁄¤‹›ﬁﬂ‡·‚„‰ÂÊÁËÈÍÎÏÌÓÔ\uf8ffÒÚÛÙıˆ˜¯˘˙˚¸˝˛ˇŁłŠšŽž¦ÐðÝýÞþ−×¹²³½¼¾₣ĞğİŞşĆćČčđ'
ACCENTS = {"ð": "²"}   # exception observée sur des plans (m²)

def _decode(t):
    """Police sans table Unicode: les glyphes sont numérotés dans l'ordre Macintosh (ex: '3/$1' = 'PLAN')."""
    out = []
    for ch in t:
        o = ord(ch)
        if ch in ACCENTS:
            out.append(ACCENTS[ch])
        elif 3 <= o < 3 + len(MAC_GLYPHS):
            out.append(MAC_GLYPHS[o - 3])
        else:
            out.append(ch)
    return "".join(out)

_TESSDATA = []

def _tessdata():
    """Dossier tessdata de Tesseract (None si absent). Sous Windows, Tesseract n'est souvent pas dans le PATH:
    on regarde aussi les dossiers d'installation habituels."""
    if not _TESSDATA:
        found = None
        try:
            found = pymupdf.get_tessdata()
        except Exception:
            cands = [os.environ.get("TESSDATA_PREFIX", "")]
            for base in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                         os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                         os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs")):
                cands.append(os.path.join(base, "Tesseract-OCR", "tessdata"))
            found = next((c for c in cands if c and os.path.isfile(os.path.join(c, "eng.traineddata"))), None)
        if found and not any(os.path.isfile(os.path.join(found, x + ".traineddata")) for x in ("fra", "eng")):
            found = None                                  # Tesseract incomplet (aucune langue installée)
        _TESSDATA.append(found)
    return _TESSDATA[0]

def _ocr_textpage(pg):
    """Page sans texte (scan, texte vectorisé): OCR si Tesseract est installé, sinon None."""
    td = _tessdata()
    if not td:
        return None
    W, H = pg.rect.width, pg.rect.height
    dpi = int(max(72, min(200, 72 * (30e6 / max(1.0, W * H)) ** 0.5)))   # max ~30 Mpx (A0 -> ~140 dpi)
    langs = [l for l in ("fra+eng", "fra", "eng")
             if all(os.path.isfile(os.path.join(td, x + ".traineddata")) for x in l.split("+"))]
    for lang in langs:
        try:
            return pg.get_textpage_ocr(flags=TXT_FLAGS, language=lang, dpi=dpi, full=True, tessdata=td)
        except Exception:
            continue
    return None

def _ocr_ok(t):
    """Ligne OCR plausible: au moins un vrai mot (>= 4 lettres) et surtout des lettres (pas du bruit de tampon/QR code)."""
    chars = [c for c in t if not c.isspace()]
    letters = [c for c in chars if c.isalpha()]
    return (len(letters) >= 6 and len(letters) >= 0.6 * len(chars)
            and any(len(w) >= 4 and w.isalpha() for w in re.findall(r"[^\W\d_]+", t)))

def ocr_title(spans, W, H):
    """Page scannée: les tailles de police de l'OCR ne sont pas fiables. Titre = lignes centrées en haut de page
    (présentation habituelle d'un titre), consécutives; à défaut, les lignes en évidence."""
    good = [s for s in spans if s["dir"][0] > 0.7 and _ocr_ok(s["text"]) and not _is_label(s["text"])
            and not BOILERPLATE.search(_norm(s["text"]))]
    cen = sorted((s for s in good if abs((s["rect"].x0 + s["rect"].x1) / 2 - W / 2) <= 0.05 * W
                  and s["rect"].y0 <= 0.45 * H and s["rect"].width < 0.95 * W), key=lambda s: s["rect"].y0)
    out = []
    for s in cen:
        if out and s["rect"].y0 - out[-1]["rect"].y1 > 3 * max(out[-1]["rect"].height, s["rect"].height):
            break                                     # grand vide: fin du bloc de titre
        out.append(s)
        if len(out) == 4:
            break
    if out:
        return " | ".join(s["text"] for s in out)[:400]
    return prominent_text(good)

def _merge_line(items, vx, vy):
    """Recolle les morceaux d'une même ligne (police/graisse différente, mot isolé) si aucun trait ne les sépare.
    vx: traits verticaux triés (x, y0, y1), vy: traits horizontaux triés (y, x0, x1)."""
    import bisect
    def crossed(a, b, d):
        if abs(d[0]) > 0.9:                         # ligne horizontale: un trait vertical dans l'intervalle ?
            g0, g1 = sorted((a["rect"].x1, b["rect"].x0) if d[0] > 0 else (b["rect"].x1, a["rect"].x0))
            c = (a["rect"].y0 + a["rect"].y1) / 2
            segs_, keys = vx
        elif abs(d[1]) > 0.9:
            g0, g1 = sorted((a["rect"].y1, b["rect"].y0) if d[1] > 0 else (b["rect"].y1, a["rect"].y0))
            c = (a["rect"].x0 + a["rect"].x1) / 2
            segs_, keys = vy
        else:
            return False
        for i in range(bisect.bisect_left(keys, g0 - 0.5), bisect.bisect_right(keys, g1 + 0.5)):
            if segs_[i][1] <= c <= segs_[i][2]:
                return True
        return False
    def along(s):
        d = s["dir"]
        return s["rect"].x0 * d[0] + s["rect"].y0 * d[1] if abs(d[0]) > 0.9 or abs(d[1]) > 0.9 else 0
    def perp(s):                                   # position de la ligne de base (centre perpendiculaire)
        d = s["dir"]
        return ((s["rect"].y0 + s["rect"].y1) / 2) if abs(d[0]) >= abs(d[1]) else ((s["rect"].x0 + s["rect"].x1) / 2)
    # morceaux écrits séparément (lettre de couleur différente...) : on les regroupe par ligne de base
    order = sorted(items, key=lambda s: (s["dir"], perp(s)))
    lines, cur = [], []
    for s in order:
        if cur and (cur[0]["dir"] != s["dir"] or abs(perp(cur[0]) - perp(s)) > 0.3 * max(cur[0]["size"], s["size"])):
            lines.append(cur); cur = []
        cur.append(s)
    if cur:
        lines.append(cur)
    res = []
    for ln in lines:
        ln.sort(key=along)
        res.extend(_merge_run(ln, crossed))
    return res

def _merge_run(items, crossed):
    out = []
    for s in items:
        p = out[-1] if out else None
        if p is not None and p["dir"] == s["dir"] and 0.75 <= s["size"] / max(p["size"], 0.1) <= 1.33:
            d = s["dir"]
            if abs(d[0]) > 0.9:
                gap = (s["rect"].x0 - p["rect"].x1) * (1 if d[0] > 0 else -1)
            elif abs(d[1]) > 0.9:
                gap = (s["rect"].y0 - p["rect"].y1) * (1 if d[1] > 0 else -1)
            else:
                gap = None
            sz = max(p["size"], s["size"])
            low = -0.8 * sz if p.get("sp_after") else -0.5 * sz   # la boîte inclut l'espace final: chevauchement normal
            if gap is not None and low <= gap <= 1.2 * sz and not crossed(p, s, d):
                n1, n2 = len(p["text"]), len(s["text"])
                space = gap > 0.1 * sz or p.get("sp_after") or s.get("sp_before")
                p["text"] = p["text"] + (" " if space else "") + s["text"]
                p["sp_after"] = s.get("sp_after")
                p["rect"] = p["rect"] | s["rect"]
                p["size"] = sz
                p["lum"] = min(p["lum"], s["lum"])
                p["bold"] = (p["bold"] if n1 >= n2 else s["bold"])
                continue
        out.append(dict(s))
    return out

_BOLD = {}

def _bold_fonts(pg):
    """Polices en gras de la page, d'après le fichier de police intégré: les PDF Outlook/Word nomment souvent
    les polices 'CIDFont+F2' sans indiquer le gras (ici 'Segoe UI Semibold'). Résultat mis en cache par fichier."""
    doc = pg.parent
    key = (doc.name, id(doc))
    cache = _BOLD.setdefault(key, {})
    if len(_BOLD) > 4:                                    # on ne garde que les derniers fichiers
        for k in list(_BOLD)[:-4]:
            del _BOLD[k]
    out = set()
    for f in pg.get_fonts(full=True):
        xref, base = f[0], f[3]
        if xref not in cache:
            bold = False
            try:
                buf = doc.extract_font(xref)[3]
                if buf:
                    ft = pymupdf.Font(fontbuffer=buf)
                    nm = ft.name.lower()
                    bold = bool(ft.is_bold) or any(w in nm for w in ("bold", "black", "heavy", "semibold", "demi"))
            except Exception:
                pass
            cache[xref] = bold
        if cache[xref]:
            out.add(base.split("+", 1)[1] if re.match(r"^[A-Z]{6}\+", base) else base)
    return out

def _page_data(pg, textpage=None):
    """Segments horizontaux/verticaux, lignes de texte et nombre de tracés, dans le repère affiché de la page."""
    M = pg.rotation_matrix
    ma, mb, mc, md, me, mf = M.a, M.b, M.c, M.d, M.e, M.f
    seen = set()
    segs = []
    rects = []
    rseen = set()
    n_items = 0

    PW, PH = pg.rect.width, pg.rect.height

    def add(x0, y0, x1, y1):
        # transformation vers le repère affiché + filtre (traits horizontaux/verticaux d'au moins 4 pt)
        X0, Y0 = ma * x0 + mc * y0 + me, mb * x0 + md * y0 + mf
        X1, Y1 = ma * x1 + mc * y1 + me, mb * x1 + md * y1 + mf
        if min(X0, X1) < -2 or max(X0, X1) > PW + 2 or min(Y0, Y1) < -2 or max(Y0, Y1) > PH + 2:
            return                     # trait qui sort de la page (pli, repère de coupe): pas un bord de cadre
        if abs(Y0 - Y1) < 1 and abs(X1 - X0) >= 4:
            key = (round(Y0), round(min(X0, X1)), round(max(X0, X1)), 0)
        elif abs(X0 - X1) < 1 and abs(Y1 - Y0) >= 4:
            key = (round(X0), round(min(Y0, Y1)), round(max(Y0, Y1)), 1)
        else:
            return
        if key not in seen:
            seen.add(key)
            segs.append((X0, Y0, X1, Y1))

    dark = []                                   # aplats foncés (texte blanc sur fond bleu...)
    for p in pg.get_cdrawings():
        f = p.get("fill")
        if f and len(f) >= 3 and f[0] * 0.3 + f[1] * 0.59 + f[2] * 0.11 < 0.5:
            fr = pymupdf.Rect(p["rect"]) * M; fr.normalize()
            if fr.width >= 20 and fr.height >= 8:
                dark.append(fr)
        for it in p["items"]:
            n_items += 1
            t = it[0]
            if t == "l":
                add(it[1][0], it[1][1], it[2][0], it[2][1])
            elif t == "re":
                r = it[1]
                x0, y0, x1, y1 = min(r[0], r[2]), min(r[1], r[3]), max(r[0], r[2]), max(r[1], r[3])
                add(x0, y0, x1, y0); add(x0, y1, x1, y1); add(x0, y0, x0, y1); add(x1, y0, x1, y1)
                if x1 - x0 >= 15 and y1 - y0 >= 15:           # cadre fermé: on le garde pour la détection par rectangle
                    ax, ay = ma * x0 + mc * y0 + me, mb * x0 + md * y0 + mf
                    bx, by = ma * x1 + mc * y1 + me, mb * x1 + md * y1 + mf
                    key = (round(min(ax, bx)), round(min(ay, by)), round(max(ax, bx)), round(max(ay, by)))
                    if key not in rseen:
                        rseen.add(key)
                        rects.append(key)
            elif t == "qu":
                q = it[1]
                try:
                    pts = [q.ul, q.ur, q.lr, q.ll]
                except AttributeError:
                    pts = [q[0], q[1], q[3], q[2]]     # tuple (ul, ur, ll, lr) -> parcours du contour
                for i in range(4):
                    add(pts[i][0], pts[i][1], pts[(i + 1) % 4][0], pts[(i + 1) % 4][1])
    raw = []
    garbled = set()
    bold_fonts = _bold_fonts(pg) if textpage is None else set()
    td = (pg.get_text("dict", textpage=textpage) if textpage is not None
          else pg.get_text("dict", flags=TXT_FLAGS))
    for b in td["blocks"]:
        for l in b.get("lines", []):
            dx, dy = l["dir"]
            d = (round(ma * dx + mc * dy, 3), round(mb * dx + md * dy, 3))
            for s in l["spans"]:
                t = s["text"]
                if any(ord(c) < 32 and c not in "\t\n\r" for c in t):
                    garbled.add(s["font"])
                raw.append((s, d))
    vx = sorted((x0, min(y0, y1), max(y0, y1)) for x0, y0, x1, y1 in segs if abs(x0 - x1) < 1)
    vy = sorted((y0, min(x0, x1), max(x0, x1)) for x0, y0, x1, y1 in segs if abs(y0 - y1) < 1)
    vx, vy = (vx, [v[0] for v in vx]), (vy, [v[0] for v in vy])
    items = []
    for s, d in raw:
        t = s["text"]
        if s["font"] in garbled:
            t = _decode(t)
        t = " ".join(t.split())
        if not t:
            continue
        r = pymupdf.Rect(s["bbox"]) * M; r.normalize()
        c = s["color"]
        lum = (((c >> 16) & 255) * 0.3 + ((c >> 8) & 255) * 0.59 + (c & 255) * 0.11) / 255
        if lum > 0.6:                           # texte clair sur aplat foncé: il est bien lisible
            cx, cy = (r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2
            if any(fr.x0 <= cx <= fr.x1 and fr.y0 <= cy <= fr.y1 for fr in dark):
                lum = 1 - lum
        bold = (bool(s["flags"] & 16) or "bold" in s["font"].lower() or "black" in s["font"].lower()
                or s["font"] in bold_fonts)
        raw_t = s["text"]
        items.append({"size": s["size"], "text": t, "rect": r, "lum": lum, "dir": d, "bold": bold,
                      "sp_before": raw_t[:1].isspace(), "sp_after": raw_t[-1:].isspace()})
    spans = _merge_line(items, vx, vy)
    return segs, spans, n_items, rects

def _rot_pt(k, x, y, W, H):
    if k == 0: return x, y
    if k == 1: return H - y, x
    if k == 2: return W - x, H - y
    return y, W - x            # k == 3

def _rot_dir(k, d):
    dx, dy = d
    return [(dx, dy), (-dy, dx), (-dx, -dy), (dy, -dx)][k]

def _merge(lines, minlen, tol=1.5, gap=4):
    """Fusionne les traits alignés (même coordonnée, intervalles qui se touchent). Garde ceux >= minlen."""
    lines = sorted(lines)
    groups, cur = [], []
    for c, a, b in lines:
        if cur and c - cur[0][0] > tol:
            groups.append(cur); cur = []
        cur.append((c, a, b))
    if cur:
        groups.append(cur)
    res = []
    for g in groups:
        c0 = sum(x[0] for x in g) / len(g)
        for _, a, b in sorted(g, key=lambda x: x[1]):
            if res and res[-1][3] == id(g) and a <= res[-1][2] + gap:
                res[-1][2] = max(res[-1][2], b)
            else:
                res.append([c0, a, b, id(g)])
    return [[c, a, b] for c, a, b, _ in res if b - a >= minlen]

def _frame(k, segs, spans, W, H):
    """Même page, tournée de k x 90° : retourne (largeur, hauteur, H, V, spans)"""
    W2, H2 = (W, H) if k in (0, 2) else (H, W)
    Hl, Vl = [], []
    for x0, y0, x1, y1 in segs:
        a = _rot_pt(k, x0, y0, W, H); b = _rot_pt(k, x1, y1, W, H)
        if abs(a[1] - b[1]) < 1 and abs(a[0] - b[0]) > 1:
            Hl.append([(a[1] + b[1]) / 2, min(a[0], b[0]), max(a[0], b[0])])
        elif abs(a[0] - b[0]) < 1 and abs(a[1] - b[1]) > 1:
            Vl.append([(a[0] + b[0]) / 2, min(a[1], b[1]), max(a[1], b[1])])
    sp = []
    for s in spans:
        p = _rot_pt(k, s["rect"].x0, s["rect"].y0, W, H); q = _rot_pt(k, s["rect"].x1, s["rect"].y1, W, H)
        r = pymupdf.Rect(min(p[0], q[0]), min(p[1], q[1]), max(p[0], q[0]), max(p[1], q[1]))
        d = _rot_dir(k, s["dir"])
        sp.append({**s, "rect": r, "dir": d})
    return W2, H2, Hl, Vl, sp

def _candidates(W, H, Hm, Vm, edge=0.03):
    """Boîtes possibles du cartouche: collées au bord droit/bas de la feuille et fermées par des traits.
    edge = tolérance (en fraction de la page) pour considérer qu'un trait touche le bord de la feuille."""
    r = 1 - edge
    cands = []
    for x, a, b in Vm:
        if 0.15 * W < x < r * W and b >= r * H - 15:
            if a <= edge * H + 15 and x > 0.4 * W:
                cands.append((pymupdf.Rect(x, 0, W, H), "A"))                # colonne pleine hauteur
            for y, c, d in Hm:
                if 0.2 * H < y < 0.95 * H and c <= x + 5 and d >= r * W - 15 and a <= y + 5:
                    cands.append((pymupdf.Rect(x, y, W, H), "A"))            # coin bas-droite
    for y, c, d in Hm:
        if 0.6 * H < y < r * H and c <= 0.1 * W and d >= r * W - 15:
            cands.append((pymupdf.Rect(0, y, W, H), "B"))                    # bandeau sur toute la largeur
    return cands

def _score_box(r, sp, W, H, min_txt, min_area, max_area=0.55):
    """Note un cadre candidat. Un vrai cartouche: texte lisible dans le bon sens, beaucoup de mots-étiquettes
    (Indice, Date, Phase, Échelle...), taille raisonnable. Retourne None si le cadre n'est pas plausible."""
    area = r.get_area() / (W * H)
    if not (min_area <= area <= max_area):
        return None
    # comparaisons simples (Rect.contains est très lent appelé des millions de fois: plannings, grilles)
    X0, Y0, X1, Y1 = r.x0, r.y0, r.x1, r.y1
    inside = [s for s in sp if X0 <= s["_r"][0] and Y0 <= s["_r"][1] and s["_r"][2] <= X1 and s["_r"][3] <= Y1]
    if len(inside) < min_txt:
        return None
    ltr = sum(1 for s in inside if s["dir"][0] > 0.7) / len(inside)       # texte lu de gauche à droite
    if ltr < 0.55:
        return None
    labels = sum(1 for s in inside if s["_lab"])
    # un cartouche contient aussi un titre: texte lisible (pas une étiquette) plus grand que le reste du cadre
    sz = sorted(s["size"] for s in inside)
    med = sz[len(sz) // 2]
    big = any(s["size"] >= 1.2 * med and s["dir"][0] > 0.7 and s["_title_ok"] for s in inside)
    return 3 * min(labels, 12) + 0.15 * min(len(inside), 60) + 5 * ltr - 0.12 * area * 100 + (4 if big else 0)

def find_cartouche(pg, segs=None, spans=None, rects=None):
    """Cherche le cartouche sur les 4 bords de la feuille (4 rotations) avec plusieurs méthodes:
      1. cadre fermé par des traits longs (colonne, coin, bandeau), tolérance stricte puis large (petits cartouches)
      2. petit rectangle fermé collé à un coin
    Tous les candidats sont notés et le meilleur est gardé. Retourne None si rien de plausible:
    on ne devine jamais une zone au hasard (sinon on lirait les grands textes du plan)."""
    if segs is None:
        segs, spans, n_items, rects = _page_data(pg)
    W0, H0 = pg.rect.width, pg.rect.height
    best = None            # (score, -aire, k, rect, frame)
    for k in (0, 1, 3, 2):
        frame = _frame(k, segs, spans, W0, H0)
        W, H, Hl, Vl, sp = frame
        Hm, Vm = _merge(Hl, 0.15 * W), _merge(Vl, 0.04 * H)
        boxes = []
        for edge, min_txt, min_area in ((0.03, 8, 0.015), (0.07, 5, 0.004)):
            for r, kind in _candidates(W, H, Hm, Vm, edge):
                boxes.append((r, min_txt, min_area))
        for (ax, ay, bx, by) in (rects or []):                          # petit rectangle collé au coin bas-droite
            p1, p2 = _rot_pt(k, ax, ay, W0, H0), _rot_pt(k, bx, by, W0, H0)
            r = pymupdf.Rect(min(p1[0], p2[0]), min(p1[1], p2[1]), max(p1[0], p2[0]), max(p1[1], p2[1]))
            if r.x1 >= 0.92 * W and r.y1 >= 0.92 * H:
                boxes.append((r, 4, 0.003))
        for s in sp:                                                    # calculés une fois par repère
            s["_r"] = (s["rect"].x0, s["rect"].y0, s["rect"].x1, s["rect"].y1)
            s["_lab"] = _is_label(s["text"]) and len(re.findall(r"[A-Za-z]", s["text"])) >= 3
            s["_title_ok"] = not _is_label(s["text"]) and _readable(s["text"])
        seen_boxes = set()
        for r, min_txt, min_area in boxes:
            kb = (round(r.x0), round(r.y0), round(r.x1), round(r.y1), min_txt, min_area)
            if kb in seen_boxes:                                        # même cadre trouvé plusieurs fois
                continue
            seen_boxes.add(kb)
            sc = _score_box(r, sp, W, H, min_txt, min_area, 0.55 if min_area >= 0.004 and min_txt >= 5 else 0.25)
            if sc is None:
                continue
            key = (sc, -r.get_area())
            if best is None or key > best[0]:
                best = (key, k, r, frame)
    if best is not None:
        return best[1], best[2], best[3]
    # 3e méthode (sans cadre): groupe de mots-étiquettes (Indice, Date, Phase, Échelle...) près d'un bord de la feuille
    best = None
    for k in (0, 1, 3, 2):
        W, H, Hl, Vl, sp = _frame(k, segs, spans, W0, H0)
        labs = [s for s in sp if s["dir"][0] > 0.7 and len(re.findall(r"[A-Za-z]", s["text"])) >= 3 and _is_label(s["text"])
                and (s["rect"].x0 > 0.6 * W or s["rect"].y0 > 0.8 * H)]
        if len(labs) < 4:
            continue
        r = pymupdf.Rect(labs[0]["rect"])
        for s_ in labs[1:]:
            r |= s_["rect"]
        for _ in range(8):                      # on agrandit le groupe avec les textes voisins (valeurs, titre)
            grown = False
            for s_ in sp:
                if s_["dir"][0] <= 0.7 or r.contains(s_["rect"].tl) and r.contains(s_["rect"].br):
                    continue
                near = pymupdf.Rect(r.x0 - 40, r.y0 - 40, r.x1 + 40, r.y1 + 40)
                if near.intersects(s_["rect"]) and (s_["rect"].x1 > 0.5 * W or s_["rect"].y0 > 0.7 * H):
                    r |= s_["rect"]
                    grown = True
            if not grown:
                break
        r = pymupdf.Rect(max(0, r.x0 - 5), max(0, r.y0 - 5), min(W, r.x1 + 5), min(H, r.y1 + 5))
        if r.get_area() / (W * H) > 0.5:
            continue
        if best is None or len(labs) > best[0]:
            best = (len(labs), k, r, (W, H, Hl, Vl, sp))
    if best is None:
        return None
    return best[1], best[2], best[3]

# ---------- Découpage du cartouche en cellules (zones fermées par les traits) ----------
def _cells(box, Hl, Vl, sp):
    bx0, by0, bx1, by1 = box
    w, h = bx1 - bx0, by1 - by0
    # traits découpés en morceaux (pointillés, export Ghostscript/PDF24): on recolle les morceaux espacés de < 8 pt
    Hl = _merge([tuple(x) for x in Hl], 0, tol=1.0, gap=8)
    Vl = _merge([tuple(x) for x in Vl], 0, tol=1.0, gap=8)
    step = max(w, h) / 450.0
    nx, ny = int(w / step) + 2, int(h / step) + 2
    blocked = [[False] * nx for _ in range(ny)]
    minlen = 0.08 * min(w, h)
    for y, a, b in Hl:
        if by0 - 2 <= y <= by1 + 2 and b - a >= minlen and b >= bx0 and a <= bx1:
            r = min(ny - 1, max(0, int((y - by0) / step)))
            c0 = max(0, int((max(a, bx0) - bx0) / step) - 1); c1 = min(nx - 1, int((min(b, bx1) - bx0) / step) + 1)
            row = blocked[r]
            for c in range(c0, c1 + 1):
                row[c] = True
    for x, a, b in Vl:
        if bx0 - 2 <= x <= bx1 + 2 and b - a >= minlen and b >= by0 and a <= by1:
            c = min(nx - 1, max(0, int((x - bx0) / step)))
            r0 = max(0, int((max(a, by0) - by0) / step) - 1); r1 = min(ny - 1, int((min(b, by1) - by0) / step) + 1)
            for r in range(r0, r1 + 1):
                blocked[r][c] = True
    label = [[0] * nx for _ in range(ny)]
    sizes = {}
    n = 0
    for r0 in range(ny):
        for c0 in range(nx):
            if blocked[r0][c0] or label[r0][c0]:
                continue
            n += 1
            stack = [(r0, c0)]
            label[r0][c0] = n
            cnt = 0
            while stack:
                r, c = stack.pop(); cnt += 1
                for rr, cc in ((r + 1, c), (r - 1, c), (r, c + 1), (r, c - 1)):
                    if 0 <= rr < ny and 0 <= cc < nx and not blocked[rr][cc] and not label[rr][cc]:
                        label[rr][cc] = n
                        stack.append((rr, cc))
            sizes[n] = cnt * step * step
    def cell_of(s):
        cx, cy = (s["rect"].x0 + s["rect"].x1) / 2, (s["rect"].y0 + s["rect"].y1) / 2
        c = min(nx - 1, max(0, int((cx - bx0) / step))); r = min(ny - 1, max(0, int((cy - by0) / step)))
        for d in range(0, 4):                       # si le centre tombe sur un trait, on cherche autour
            for rr in range(max(0, r - d), min(ny, r + d + 1)):
                for cc in range(max(0, c - d), min(nx, c + d + 1)):
                    if label[rr][cc]:
                        return label[rr][cc]
        return 0
    return sizes, cell_of

ANCHOR_RATIO = 0.76
CHAIN_RATIO = 0.38      # 1re ligne sous le titre (sous-titre): >= 38% de la taille du titre
CHAIN_RATIO_NEXT = 0.7   # lignes suivantes: proches de la taille de la ligne au-dessus (évite les étiquettes)
MAX_CHAIN = 8        # nombre max de lignes ajoutées autour du titre
CHAIN_GAP = 3.5       # écart vertical max (en hauteurs de ligne) entre deux lignes du titre
BIG_CELL = 0.3     # on ne garde que les cellules >= 30% de la plus grande cellule contenant du texte

LABEL_WORDS = set("""batiment ouvrage zone discipline specialite processus titre titres phase indice date echelle niveau
sequence format projet type spec emetteur lot volet entreprise libelle auteur dessinateur verificateur approbateur
maitrise oeuvre design manager architecte chef construction bureaux controle coordonnateur sps opc et de d l par emis
information complementaire ordre n valeur revision folio""".split())

def _is_label(text):
    t = unicodedata.normalize("NFKD", text.lower())
    t = "".join(c for c in t if not unicodedata.combining(c))
    words = re.findall(r"[a-z]+", t)
    return bool(words) and all(w in LABEL_WORDS for w in words)

DOC_MAX_ITEMS = 1500   # au-delà, la page est un plan (beaucoup de traits): jamais de texte central

BOILERPLATE = re.compile(
    r"propriete d|ne peut etre ni reproduit|reproduction interdite|ni diffuse|sans autorisation|tous droits reserves|"
    r"confidentiel|direction engineering|architecture, design and technical|groupe adp|^page \d|^\d+ ?/ ?\d+$|copyright|©")

# Étiquette qui annonce le titre dans le cartouche ("Titre :", "Désignation", "Intitulé du plan"...)
TITLE_LABEL = re.compile(r"^(?:titres?|intitule|designation|libelle|nom(?= d))(?: du| de la| de l'| de)?"
                         r"(?: plan| document| dessin| dossier)?\b\s*[:\-.]?\s*(.*)$")
STOP = set("de du des le la les l d et en au aux a pour sur sous par un une avec".split())

def _tokens(text):
    return {w for w in re.findall(r"[a-z0-9]+", _norm(text)) if len(w) >= 3 and w not in STOP}

def _overlap(hint_tokens, text):
    """Part des mots du nom de fichier retrouvés dans le texte (comparaison sur les 4 premières lettres: abréviations)."""
    if not hint_tokens:
        return 0.0
    tt = {w[:4] for w in _tokens(text)}
    return sum(1 for w in hint_tokens if w[:4] in tt) / len(hint_tokens)

def _norm(text):
    t = unicodedata.normalize("NFKD", text.lower())
    return "".join(c for c in t if not unicodedata.combining(c))

def _body_size(spans):
    """Taille du corps de texte = médiane des tailles des lignes (un long titre ne fausse pas le calcul)."""
    sz = sorted((sp["size"], 1) for sp in spans)
    if not sz:
        return 0
    half, acc = sum(n for _, n in sz) / 2, 0
    for v, n in sz:
        acc += n
        if acc >= half:
            return v
    return sz[-1][0]

def prominent_text(spans, ratio=0.7, n=10):
    """Toutes les lignes en évidence de la page (plus grosse police, en gras un peu plus petit), dans l'ordre de lecture.
    Sert de repli quand aucun titre n'est trouvé autrement."""
    return _join(prominent_spans(spans, ratio, n))

def _join(spans):
    out = []
    for sp in spans:
        if sp["text"] not in out:
            out.append(sp["text"])
    return " | ".join(out)[:600]

def prominent_spans(spans, ratio=0.7, n=10):
    c = []
    body = _body_size(spans)
    for sp in spans:
        t = sp["text"]
        legal = BOILERPLATE.search(_norm(t)) and (sp["size"] < 1.5 * body or len(t) > 60)   # "Groupe ADP" en gros titre: gardé
        if len(re.findall(r"[A-Za-zÀ-ÿ]", t)) < 2 or legal or _is_label(t) or sp["lum"] > 0.85:
            continue
        if (" " not in t and re.search(r"\d", t)) or _is_code(t):
            continue
        c.append(sp)
    if not c:
        return []
    # sens de lecture du plus gros texte (page tournée: texte vertical), on garde les textes dans ce sens
    dx, dy = max(c, key=lambda sp: sp["size"])["dir"]
    c = [sp for sp in c if abs(sp["dir"][0] - dx) < 0.3 and abs(sp["dir"][1] - dy) < 0.3]
    mx = max(sp["size"] for sp in c)
    top = [sp for sp in c if sp["size"] >= ratio * mx or (sp["bold"] and sp["size"] >= (ratio - 0.1) * mx)]
    top = sorted(top, key=lambda sp: -sp["size"])[:n]
    def reading(sp):                                  # ordre de lecture dans le repère du texte
        cx, cy = (sp["rect"].x0 + sp["rect"].x1) / 2, (sp["rect"].y0 + sp["rect"].y1) / 2
        return (round((-dy * cx + dx * cy) / (mx * 0.5)), cx * dx + cy * dy)
    top.sort(key=reading)
    return top

COVER_RATIO = 1.8   # document: police du titre >= 1,8 x corps du texte -> page de garde, on prend les textes en évidence

def central_text(pg):
    """Documents sans cartouche (Word, présentation): ensemble des lignes en évidence au centre de la page.
    En évidence = plus grosses que le corps du texte et/ou en gras. Les mentions légales / pieds de page sont ignorés."""
    M = pg.rotation_matrix
    W, H = pg.rect.width, pg.rect.height
    blocks, sizes = [], []
    for b in pg.get_text("dict", flags=TXT_FLAGS)["blocks"]:
        if b.get("type") != 0:
            continue
        lines, size, nb, nchar, row, lrects = [], 0.0, 0, 0, False, []
        for l in b["lines"]:
            dx, dy = l["dir"]
            if abs(M.b * dx + M.d * dy) > 0.1:
                continue
            raw = "".join(sp["text"] for sp in l["spans"])
            t = " ".join(raw.split())
            if not t:
                continue
            lines.append(t)
            lr = pymupdf.Rect()
            for sp in l["spans"]:
                if sp["text"].strip():
                    rr = pymupdf.Rect(sp["bbox"]) * M; rr.normalize()
                    lr |= rr if lr.is_empty else rr
            lrects.append((pymupdf.Rect(l["bbox"]) * M, max((sp["size"] for sp in l["spans"]), default=8)))
            if re.search(r"\s{3,}", raw.strip()):
                row = True
            sps = sorted((pymupdf.Rect(sp["bbox"]) * M for sp in l["spans"] if sp["text"].strip()), key=lambda r_: r_.x0)
            for r1, r2 in zip(sps, sps[1:]):
                if r2.x0 - r1.x1 > 1.5 * max(l["spans"][0]["size"], 6):
                    row = True                     # grand vide au milieu de la ligne = cellules de tableau
            for sp in l["spans"]:
                n = len(sp["text"].strip())
                if n:
                    size = max(size, sp["size"])
                    sizes.append((sp["size"], n))
                nchar += n
                fn_ = sp["font"].lower()
                if (sp["flags"] & 16) or any(w in fn_ for w in ("bold", "black", "demi", "heavy")):
                    nb += n
        for i1 in range(len(lrects)):               # lignes côte à côte dans le bloc = cellules de tableau
            for i2 in range(i1 + 1, len(lrects)):
                (ra, sa), (rb, sb) = lrects[i1], lrects[i2]
                ra.normalize(); rb.normalize()
                ov = min(ra.y1, rb.y1) - max(ra.y0, rb.y0)
                gap = max(rb.x0 - ra.x1, ra.x0 - rb.x1)
                if ov > 0.5 * min(ra.height, rb.height) and gap > 1.5 * max(sa, sb, 6):
                    row = True
        text = " ".join(lines)
        if len(re.findall(r"[A-Za-zÀ-ÿ]", text)) < 3:
            continue
        r = pymupdf.Rect(b["bbox"]) * M; r.normalize()
        blocks.append({"text": text, "size": size, "bold": nb / max(1, nchar) >= 0.5, "rect": r, "row": row})
    if not blocks:
        return ""
    # taille du corps de texte = médiane pondérée par le nombre de caractères
    sizes.sort()
    half, acc, body = sum(n for _, n in sizes) / 2, 0, sizes[-1][0]
    for sz, n in sizes:
        acc += n
        if acc >= half:
            body = sz
            break
    cand = [b for b in blocks
            if 0.05 * H < (b["rect"].y0 + b["rect"].y1) / 2 < 0.95 * H and not b["row"] and not BOILERPLATE.search(_norm(b["text"]))]
    if not cand:
        return ""
    # lignes en évidence, dans la partie centrale de la page
    def evidence(b):
        cx = (b["rect"].x0 + b["rect"].x1) / 2
        centre = abs(cx - W / 2) <= 0.06 * W and b["rect"].width <= 0.8 * W      # ligne centrée (titre), pas un paragraphe justifié
        return (b["size"] >= 1.12 * body                                         # plus grand que le corps du texte
                or (b["bold"] and b["size"] >= 0.6 * body)                       # en gras, même si la police est plus petite
                or (centre and b["size"] >= 0.9 * body))
    prom = [b for b in cand
            if (0.18 * H <= (b["rect"].y0 + b["rect"].y1) / 2 <= 0.82 * H and evidence(b))
            or b["size"] >= 1.3 * body]                # texte nettement plus grand que le corps: retenu même hors de la zone centrale
    # on retire les lignes de tableau (plusieurs blocs en évidence sur la même ligne)
    def same_row(a_, b_):
        ov = min(a_["rect"].y1, b_["rect"].y1) - max(a_["rect"].y0, b_["rect"].y0)
        return ov > 0.5 * min(a_["rect"].height, b_["rect"].height) and (a_["rect"].x1 < b_["rect"].x0 or b_["rect"].x1 < a_["rect"].x0)
    prom = [b for b in prom if not any(o is not b and same_row(b, o) for o in prom)]
    if not prom:                              # repli: le bloc le plus en évidence, le plus proche du centre
        def score(b):
            cx, cy = (b["rect"].x0 + b["rect"].x1) / 2, (b["rect"].y0 + b["rect"].y1) / 2
            return b["size"] * (1.3 if b["bold"] else 1.0) * (1 - 0.35 * abs(cx - W / 2) / (W / 2) - 0.35 * abs(cy - H / 2) / (H / 2))
        marked = [b for b in cand if b["bold"] or b["size"] >= 1.05 * body]
        prom = sorted(marked or cand, key=score, reverse=True)[:3 if marked else 1]
    for _ in range(3):                        # on ajoute les blocs collés à un bloc en évidence (sous-titre, suite)
        added = None
        for b in cand:
            if b in prom:
                continue
            for k in prom:
                gap = max(b["rect"].y0 - k["rect"].y1, k["rect"].y0 - b["rect"].y1)
                ov = min(b["rect"].x1, k["rect"].x1) - max(b["rect"].x0, k["rect"].x0)
                if gap <= 1.3 * k["size"] and ov > 0 and 0.5 * k["size"] <= b["size"] <= 1.1 * k["size"]:
                    added = b
                    break
            if added:
                break
        if not added:
            break
        prom.append(added)
    prom.sort(key=lambda b: (b["rect"].y0, b["rect"].x0))
    out = []
    for b in prom:
        if b["text"] not in out:
            out.append(b["text"])
    return " | ".join(out)[:800]

def _assemble(ok, chain=True):
    """Titre = lignes en plus grande police (ancres) + lignes voisines (au-dessus et en dessous) qui forment le même
    bloc: écart vertical <= CHAIN_GAP hauteurs de ligne (1 à 2 lignes vides tolérées), alignées horizontalement."""
    if not ok:
        return ""
    mx = max(s["size"] for s in ok)
    for s_ in ok:
        s_["anchor"] = s_["size"] >= mx * ANCHOR_RATIO
    keep = [s_ for s_ in ok if s_["anchor"]]
    # On ajoute les lignes qui prolongent le titre (projet au-dessus, sous-titre, zone en dessous...).
    # Chaque ligne est comparée à sa voisine déjà retenue (pas à la plus grande police du cartouche).
    for _ in range(MAX_CHAIN if chain else 0):
        best = None
        for c in ok:
            if c in keep:
                continue
            for kk in keep:
                gap = max(c["rect"].y0 - kk["rect"].y1, kk["rect"].y0 - c["rect"].y1)   # dessous ou dessus
                overlap = min(c["rect"].x1, kk["rect"].x1 + 0.1 * kk["rect"].width) - \
                          max(c["rect"].x0, kk["rect"].x0 - 0.1 * kk["rect"].width)
                if (-0.3 * kk["size"] <= gap <= CHAIN_GAP * kk["size"] and overlap > 0
                        and (CHAIN_RATIO if kk["anchor"] else CHAIN_RATIO_NEXT) * kk["size"] <= c["size"] <= kk["size"] * 1.05):
                    if best is None or gap < best[0]:
                        best = (gap, c)
        if best is None:
            break
        best[1]["anchor"] = False
        keep.append(best[1])
    keep.sort(key=lambda s: (round(s["rect"].y0 / (mx * 0.4)), s["rect"].x0))
    return _join(keep)

def _is_code(t):
    """Ligne de codification ('CDGCIPFR23_00 1190 DCE G CAL GEN XTN001'): presque que des codes courts ou chiffrés."""
    w = t.split()
    if len(w) < 4:
        return False
    return sum(1 for x in w if re.search(r"[\d_]", x) or (len(x) <= 3 and x.isupper()) or not re.search(r"[A-Za-z]", x)) >= 0.7 * len(w)

def _readable(t):
    """Texte qui peut faire partie d'un titre (pas un code, une date, un indice...)."""
    if len(re.findall(r"[A-Za-zÀ-ÿ]", t)) < 4:
        return False
    if " " not in t and re.search(r"\d", t):
        return False
    return not _is_code(t) and not BOILERPLATE.search(_norm(t))

def _strip_label(t):
    """'Titre : PLAN DU RDC' -> 'PLAN DU RDC'. Retourne (texte, True si une étiquette de titre a été trouvée)."""
    n = _norm(t)
    m = TITLE_LABEL.match(n)
    if not m:
        return t, False
    rest = t[m.start(1):] if len(n) == len(t) else n[m.start(1):].upper()
    return rest.strip(" :-."), True

def _by_label(ok_all, labels, cell_of):
    """Texte désigné par une étiquette 'Titre' / 'Désignation': même cellule, sinon la cellule juste à droite ou dessous."""
    for lab in sorted(labels, key=lambda s: -s["size"]):
        lr = lab["rect"]
        same = [s for s in ok_all if s is not lab and s["cell"] == lab["cell"] and s["cell"]
                and s["rect"].y0 >= lr.y0 - 0.5 * lab["size"] and s["size"] >= 0.8 * lab["size"]]
        if not same:
            near = []
            for s in ok_all:
                if s is lab or s["size"] < 0.8 * lab["size"]:
                    continue
                r = s["rect"]
                below = r.y0 >= lr.y1 - 2 and r.x1 > lr.x0 - 5 and r.y0 - lr.y1 < 6 * lab["size"]
                right = r.x0 >= lr.x1 - 2 and r.y1 > lr.y0 and r.y0 < lr.y1 + 0.5 * lab["size"]
                if below or right:
                    near.append((max(0, r.y0 - lr.y1) + max(0, r.x0 - lr.x1), s))
            if near:
                c = min(near, key=lambda x: x[0])[1]["cell"]
                same = [s for s in ok_all if s["cell"] == c] if c else [min(near, key=lambda x: x[0])[1]]
        if same:
            t = _assemble(same)
            if t:
                return t
    return ""

def _heading_spans(spans, H, body, n=2):
    """Lignes de titre en gras en haut de page (objet d'un mail, titre Word...): courtes, pas un paragraphe,
    seules sur leur ligne (pas un en-tête de colonne de tableau: 'Prix total | Quantité | P.U.')."""
    def in_row(s):
        r = s["rect"]
        return any(o is not s and min(r.y1, o["rect"].y1) - max(r.y0, o["rect"].y0) > 0.5 * min(r.height, o["rect"].height)
                   and (o["rect"].x0 > r.x1 + 2 or o["rect"].x1 < r.x0 - 2) for o in spans)
    out = []
    for s in sorted(spans, key=lambda s: s["rect"].y0):
        t = s["text"]
        if not s["bold"] or s["dir"][0] < 0.7 or (s["rect"].y0 + s["rect"].y1) / 2 > 0.4 * H or s["size"] < 0.8 * body:
            continue
        words = [w for w in t.split() if len(re.findall(r"[A-Za-zÀ-ÿ]", w)) >= 4]
        if (not words or len(re.findall(r"[A-Za-zÀ-ÿ]", t)) < 6 or len(t) > 150 or _is_label(t) or _is_code(t)
                or BOILERPLATE.search(_norm(t)) or in_row(s)):
            continue
        out.append(s)
        if len(out) == n:
            break
    return out

def _doc_title(pg, spans, ocr="", doc=True):
    """Page sans cartouche: document (page de garde, note, mail...) ou repli sur les textes en évidence."""
    if doc and not ocr and spans:
        # Titre = lignes en gras en haut de page + lignes plus grosses que le corps du texte (les deux:
        # le titre n'est pas toujours en gras, ni toujours le plus gros)
        body = _body_size(spans)
        H = pg.rect.height
        heads = _heading_spans(spans, H, body)
        # lignes plus grosses que le corps: en haut de page (sinon titres de section d'un tableau), ou très grosses;
        # si un titre en gras existe, seulement celles plus grosses que lui (titre non gras plus gros)
        big = [s for s in prominent_spans(spans) if s["size"] >= 1.15 * body
               and ((s["rect"].y0 + s["rect"].y1) / 2 <= 0.4 * H or s["size"] >= COVER_RATIO * body)]
        if heads:
            hmax = max(s["size"] for s in heads)
            big = [s for s in big if s["size"] > 1.05 * hmax]
        if heads or big:
            both = sorted({id(s): s for s in heads + big}.values(), key=lambda s: (s["rect"].y0, s["rect"].x0))
            return _join(both), None, "texte en évidence"
    if doc and not ocr:                               # document: pas de cartouche
        body = _body_size(spans)
        if spans and max(s["size"] for s in spans) >= COVER_RATIO * body:
            t = prominent_text(spans)                 # page de garde: gros titre au milieu
            if t:
                return t, None, "texte en évidence"
        t = central_text(pg)                          # sinon paragraphe central
        if t:
            return t, None, "texte central"
    t = prominent_text(spans)
    if t:
        return t, None, "texte en évidence (à vérifier)" + ocr
    if not spans:
        return "(aucun texte - page scannée ou texte vectorisé ?)", None, "aucun texte"
    return "(cartouche non détecté)", None, "non détecté"

def cartouche_title(pg, hint=""):
    """Retourne (titre, cadre du cartouche, origine). hint = description tirée du nom de fichier (aide au choix)."""
    if not _tessdata() and (not pg.get_fonts(full=True) or not pg.get_text("text", flags=TXT_FLAGS).strip()):
        # aucun texte et pas d'OCR possible: inutile d'analyser les traits (très long sur un plan vectorisé)
        return "(aucun texte - page scannée ou texte vectorisé ?)", None, "aucun texte"
    segs, spans, n_items, rects = _page_data(pg)
    ocr = ""
    if not spans:                                         # aucun texte: scan ou texte vectorisé -> OCR si possible
        tp = _ocr_textpage(pg)
        if tp is not None:
            segs, spans, n_items, rects = _page_data(pg, tp)
            ocr = " (OCR)"
    found = find_cartouche(pg, segs, spans, rects)
    doc = n_items <= DOC_MAX_ITEMS
    if found is None and ocr:
        t = ocr_title(spans, pg.rect.width, pg.rect.height)
        if t:
            return t, None, "texte OCR (à vérifier)"
    if found is None:
        return _doc_title(pg, spans, ocr, doc)
    k, box, (W, H, Hl, Vl, sp) = found
    # Le cadre trouvé n'est pas toujours un cartouche: tableau des révisions en bas d'une page de garde,
    # bloc de lignes d'un planning (Gantt)... Si un texte bien plus gros existe hors du cadre, c'est lui le titre.
    inside = [s for s in sp if box.contains(s["rect"].tl) and box.contains(s["rect"].br)]
    outside = [s for s in sp if not (box.contains(s["rect"].tl) and box.contains(s["rect"].br))]
    mx_in = max((s["size"] for s in inside), default=0)
    cover = ""
    n_lab = sum(1 for s in inside if _is_label(s["text"]) and len(re.findall(r"[A-Za-z]", s["text"])) >= 3)
    t = prominent_text(outside)
    if t:
        mx_out = max(s["size"] for s in outside if s["text"] in t.split(" | "))
        weak = n_lab < 2                                  # aucun mot d'étiquette (Indice, Date, Phase...): pas un cartouche
        if mx_out >= 1.5 * mx_in and (weak or (doc and box.width >= 0.8 * W)):
            return t, None, "texte en évidence"           # planning, tableau des révisions en bas de page de garde
        if mx_out >= 1.5 * mx_in and doc and mx_out >= 2.5 * mx_in:
            cover = t                                     # gros titre de page de garde + vrai cartouche: on garde les deux
    long_ratio = sum(1 for s in inside if len(s["text"]) >= 50) / max(1, len(inside))
    if (doc and n_lab < 3) or long_ratio >= 0.15:
        # pas un cartouche: signature / mentions d'un mail, colonne "Commentaires" d'un tableau (phrases longues)
        found = None
    if found is None:
        return _doc_title(pg, spans, ocr, doc)
    ok_all, labels = [], []
    for s in sp:
        r = s["rect"]
        if not (box.contains(r.tl) and box.contains(r.br)):
            continue
        if abs(s["dir"][1]) > 0.1 or s["lum"] > 0.6:
            continue
        t, is_lab = _strip_label(s["text"])
        if is_lab:
            labels.append(s)
            if not _readable(t):
                continue                                  # étiquette seule: elle sert à localiser le titre
            s = {**s, "text": t}
        elif _is_label(s["text"]) or not _readable(s["text"]):
            continue                                      # étiquette du cartouche ("Phase", "Indice"...), code, mention légale
        ok_all.append(s)
    if not ok_all:
        t = prominent_text([s for s in sp if box.contains(s["rect"].tl) and box.contains(s["rect"].br)]) or prominent_text(sp)
        return (t, box, "texte en évidence (à vérifier)" + ocr) if t else ("", box, "cartouche" + ocr)
    sizes, cell_of = _cells(box, Hl, Vl, sp)
    for s in ok_all + labels:
        s["cell"] = cell_of(s)
    # texte foncé de préférence (le gris clair est souvent un filigrane ou un texte de gabarit)
    ok = [s for s in ok_all if s["lum"] <= 0.4] or ok_all
    amax = max(sizes.get(s["cell"], 0) for s in ok)
    if amax > 0:
        ok = [s for s in ok if sizes.get(s["cell"], 0) >= BIG_CELL * amax] or ok
    choices = [(_assemble(ok), "cartouche")]
    t_lab = _by_label(ok_all, labels, cell_of) if labels else ""
    if t_lab:
        choices.insert(0, (t_lab, "cartouche (étiquette Titre)"))
    hint_t = _tokens(hint)
    best_t, best_o = choices[0]
    if hint_t:
        scored = [(_overlap(hint_t, t), -i, t, o) for i, (t, o) in enumerate(choices)]
        sc, _, best_t, best_o = max(scored)
        if sc < 0.5:                                      # aucun choix ne ressemble au nom de fichier: on cherche la cellule qui y ressemble
            cells = {}
            for s in ok_all:
                cells.setdefault(s["cell"], []).append(s)
            alt = max(((_overlap(hint_t, _assemble(v)), _assemble(v)) for c, v in cells.items() if c), default=(0, ""))
            if alt[0] >= 0.5 and alt[0] > sc and len(hint_t) >= 2:
                best_t, best_o = alt[1], "cartouche (proche nom fichier)"
    # Titre complet = toutes les lignes de la cellule du titre (aéroport, projet, type de plan, objet...):
    # plus de mots-clés pour la recherche. Seulement si la cellule est bien délimitée (pas tout le cartouche).
    parts = best_t.split(" | ")
    cells = [s["cell"] for s in ok_all if s["text"] in parts and s["cell"]]
    if cells:
        c = max(set(cells), key=cells.count)
        if sizes.get(c, 0) < 0.6 * box.get_area():
            full = [s for s in ok_all if s["cell"] == c]
            mx_c = max(s["size"] for s in full)
            full.sort(key=lambda s: (round(s["rect"].y0 / (mx_c * 0.4)), s["rect"].x0))
            t_full = _join(full)
            if len(t_full) > len(best_t):
                best_t = t_full
    if not best_t.strip():
        t = prominent_text(sp)
        return (t, box, "texte en évidence (à vérifier)" + ocr) if t else ("", box, "cartouche" + ocr)
    if cover:
        parts = cover.split(" | ")
        parts += [x for x in best_t.split(" | ") if x not in parts]
        best_t, best_o = " | ".join(parts)[:600], "texte en évidence + cartouche"
    return best_t, box, best_o + ocr

def debug(path, page=0):
    """python inventaire.py --debug fichier.pdf [n° page]: affiche ce que le script voit sur la page."""
    d = pymupdf.open(path)
    pg = d[page]
    segs, spans, n_items, rects = _page_data(pg)
    print(f"Page {page + 1}/{len(d)}  taille {pg.rect.width:.0f}x{pg.rect.height:.0f}  rotation {pg.rotation}  "
          f"tracés {n_items}  lignes de texte {len(spans)}  corps {_body_size(spans):.1f}")
    for s in sorted(spans, key=lambda s: -s["size"])[:40]:
        print(f"  {s['size']:6.1f}  lum {s['lum']:.2f}  {'G' if s['bold'] else ' '}  dir {s['dir']}  "
              f"({s['rect'].x0:.0f},{s['rect'].y0:.0f})  {s['text'][:80]}")
    desc = parse_name(pathlib.Path(path).stem)[1 + len(CODE_FIELDS)]
    t, box, o = cartouche_title(pg, desc)
    if o.startswith("aucun texte"):
        print(f"\nAucun texte sur la page ({len(pg.get_fonts())} police, {len(pg.get_images())} image): "
              + ("scan (image)" if pg.get_images() else "texte transformé en traits (vectorisé)"))
        t, o = fallback_title(d.metadata.get("title", ""), pathlib.Path(path).stem, desc, page_kind(pg))
    print(f"\nCadre: {box}\nOrigine: {o}\nTitre: {t}")

GUI = len(sys.argv) < 2
STAMP = datetime.datetime.now().strftime("%Y-%m-%d_%Hh%M")

# ---- Codification du nom de fichier (à adapter si besoin) ----
# Convention 1 (ADP): CDGCIPFR23_001_1120_DCE_A_CD_ESC_7TN910_Circulations_verticales_Nodal_A
#   Projet = CDGCIPFR23_001 | Bâtiment = 1120 | Phase = DCE | Discipline = A
#   Type = CD | Spec = ESC | N° d'ordre = 7TN910 | Description = le reste | Indice = dernier bloc
# Convention 2 (Eiffage): 59NLIL1-8048-PRO-STR-EBES-PLA_ABCD-SS2-002-A-PLANCHER_HAUT_SS2
#   Projet = 59NLIL1-8048 | Phase = PRO | Lot = STR | Émetteur = EBES | Type = PLA | Bât. = ABCD
#   Niveau = SS2 | N° = 002 | Indice = A | Description = PLANCHER HAUT SS2
CODE_FIELDS = ["Bâtiment", "Phase", "Discipline / Lot", "Type", "Spec / Niveau", "N° d'ordre"]
N_PROJET = 2  # nombre de blocs formant le "Projet" (convention 1)
RE_EIFFAGE = re.compile(r"^([^-_]+)-(\d+)-([A-Z]{3})-([A-Z]{2,5})-([A-Z0-9]+)-([A-Z]{3})(?:_([A-Z]+))?"
                        r"-([A-Z0-9]+)-(\d+)-([A-Z])(?:-(.*))?$")

def parse_name(stem):
    m = RE_EIFFAGE.match(stem)
    if m:
        g = m.groups()
        desc = " ".join(x for x in [g[4], (g[10] or "").replace("_", " ")] if x)
        return [f"{g[0]}-{g[1]}", g[6] or "", g[2], g[3], g[5], g[7], g[8], desc, g[9]]
    t = stem.split("_")
    need = N_PROJET + len(CODE_FIELDS)
    if len(t) < need:
        return [""] * (1 + len(CODE_FIELDS)) + [stem.replace("_", " "), ""]
    projet = "_".join(t[:N_PROJET])
    codes = t[N_PROJET:need]
    rest = t[need:]
    indice = rest.pop() if rest and len(rest[-1]) <= 2 else ""
    return [projet] + codes + [" ".join(rest), indice]

# Inventaire: colonnes texte (copiables) d'abord, colonnes calculées (formules) à la fin.
# Filtrer "Trouvé" = 1 donne les résultats de la feuille Recherche en texte simple.
HEADERS = (["Titre détecté", "Chemin", "Lien", "Nom", "Sous-dossier", "Projet"] + CODE_FIELDS +
           ["Description (nom fichier)", "Indice", "Pages", "Taille (Mo)", "Modifié le", "Titre PDF", "Origine titre",
            "Trouvé", "Lien relatif", "Index", "N° trouvé"])
C_LIEN = HEADERS.index("Lien") + 1
C_CHEMIN = HEADERS.index("Chemin") + 1
C_REL = HEADERS.index("Lien relatif") + 1   # chemin relatif au fichier Excel (liens courts: HYPERLINK limité à 255 car.)
C_INDEX = HEADERS.index("Index") + 1
C_MATCH = HEADERS.index("Trouvé") + 1        # 1 si la ligne contient tous les mots-clés de Recherche!B1
C_CNT = HEADERS.index("N° trouvé") + 1       # compteur cumulé des lignes trouvées (sert à les lister dans Recherche)
C_TEXTE = HEADERS.index("Titre détecté") + 1
WIDTHS = {"Nom": 45, "Sous-dossier": 30, "Description (nom fichier)": 40,
          "Titre PDF": 25, "Titre détecté": 60, "Origine titre": 14, "Chemin": 40, "Lien": 9, "Index": 8,
          "Ouvrir": 9, "Lien relatif": 30, "Trouvé": 9}
# Ordre des colonnes de la feuille Recherche: ce qui sert à retrouver le fichier d'abord
# "Ouvrir" = lien calculé, les autres = colonnes de l'Inventaire
SEARCH_COLS = (["Titre détecté", "Chemin", "Ouvrir", "Nom", "Sous-dossier", "Projet"] + CODE_FIELDS +
               ["Description (nom fichier)", "Indice", "Pages", "Taille (Mo)", "Modifié le", "Titre PDF",
                "Origine titre"])

def list_pdfs(root):
    out = []
    for dirpath, _, files in os.walk(root):
        for fn in files:
            if fn.lower().endswith(".pdf"):
                out.append(pathlib.Path(dirpath) / fn)
    out.sort()
    return out

# ---- Recherche par mots-clés: formules classiques (aucune formule matricielle, marche dans tout Excel) ----
KW = "Recherche!$B$1"
MAX_WORDS = 8
# les mots de B1, un par élément (vides au-delà du nombre de mots): découpage par SUBSTITUE/STXT
WORDS = (f'TRIM(MID(SUBSTITUTE(TRIM({KW})," ",REPT(" ",200)),'
         f'{{{",".join(str(i * 200 + 1) for i in range(MAX_WORDS))}}},200))')

def match_formula(cell):
    """1 si la cellule contient tous les mots-clés (ordre libre), ou si la recherche est vide."""
    return (f'=IF(TRIM({KW})="",1,--(SUMPRODUCT(({WORDS}<>"")*ISNUMBER(SEARCH({WORDS},{cell})))'
            f'=SUMPRODUCT(--({WORDS}<>""))))')

def highlight_formula(cell):
    """Mise en forme conditionnelle: la cellule contient au moins un des mots-clés."""
    return f'AND(TRIM({KW})<>"",SUMPRODUCT(({WORDS}<>"")*ISNUMBER(SEARCH({WORDS},{cell})))>0)'

def page_kind(pg):
    """Pourquoi une page n'a pas de texte: image scannée, texte transformé en traits, ou page vide."""
    if pg.get_images():
        return "scan"
    return "texte vectorisé" if len(pg.read_contents()) > 20000 else "page vide"

def fallback_title(meta, stem, desc, kind="scan"):
    """PDF sans texte (scan, ou texte transformé en traits) et pas d'OCR: titre des propriétés du PDF
    s'il a un sens, sinon le nom du fichier. kind = "scan" ou "texte vectorisé" (indiqué dans l'origine)."""
    t = (meta or "").strip().strip("()[]\"'« »").strip()
    t = re.sub(r"^(microsoft (word|excel|powerpoint)|word|excel)\s*-\s*", "", t, flags=re.I)
    t = re.sub(r"\.(docx?|xlsx?|pptx?|pdf|dwg)$", "", t, flags=re.I).strip()
    same = _norm(t).replace("_", " ").strip() in (_norm(stem).replace("_", " "), "")
    words = [w for w in t.split() if len(re.findall(r"[A-Za-zÀ-ÿ]", w)) >= 3]
    if t and not same and len(words) >= 2 and not _is_code(t):     # vraie phrase, pas "scan_ADP_001"
        return t, f"propriétés PDF ({kind})"
    name = desc or stem
    return " ".join(name.replace("_", " ").split()), f"nom du fichier ({kind})"

def search_index(t):
    """Texte cherché par la feuille Recherche, + copie sans accents: 'controle' trouve aussi 'Contrôle'."""
    plain = "".join(c for c in unicodedata.normalize("NFKD", t) if not unicodedata.combining(c))
    return t if plain == t else f"{t} {plain}"

def scan(root, pdfs, state, cancel):
    """Analyse les PDF un par un et remplit le classeur Excel. state est lu par la fenêtre de progression."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventaire"
    ws.append(HEADERS)
    prepare_shortcuts(root)
    n = 0
    for p in pdfs:
        if cancel.is_set():
            break
        state["name"] = p.name
        st = p.stat()
        pages, title, txt, origine = "", "", "", ""
        parsed = parse_name(p.stem)
        try:
            d = pymupdf.open(p)
            pages = len(d)
            title = d.metadata.get("title", "") or ""
            hint = parsed[1 + len(CODE_FIELDS)]           # description tirée du nom de fichier
            txt, _box, origine = cartouche_title(d[0], hint) if pages else ("", None, "")
            if pages > 1 and (not txt or origine in ("non détecté", "aucun texte")):
                t2, _b2, o2 = cartouche_title(d[1], hint)  # page de garde vide ou sans cartouche: on essaie la page 2
                if t2 and o2 not in ("non détecté", "aucun texte"):
                    txt, origine = t2, o2 + " p.2"
            if origine.startswith("aucun texte") or not txt:
                # pas d'image sur la page = le texte a été transformé en traits (impression PDF24 / Ghostscript)
                kind = page_kind(d[0])
                txt, origine = fallback_title(title, p.stem, parsed[1 + len(CODE_FIELDS)], kind)
            d.close()
        except Exception as e:
            txt = f"ERREUR: {e}"
        rel, label = link_for(p, root, n + 1)
        vals = dict(zip(["Projet"] + CODE_FIELDS + ["Description (nom fichier)", "Indice"], parsed))
        vals.update({"Titre détecté": txt, "Chemin": str(p), "Nom": p.name,
                     "Sous-dossier": str(p.parent.relative_to(root)), "Pages": pages,
                     "Taille (Mo)": round(st.st_size / 1048576, 2),
                     "Modifié le": datetime.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d"),
                     "Titre PDF": title, "Origine titre": origine, "Lien relatif": rel,
                     "Index": search_index(f"{p.stem.replace('_', ' ')} {title} {txt}")})
        ws.append([clean(vals.get(h)) for h in HEADERS])
        r = ws.max_row
        ws.cell(row=r, column=C_MATCH, value=match_formula(f"{L(C_INDEX)}{r}"))
        ws.cell(row=r, column=C_CNT, value=f"={L(C_MATCH)}{r}" if r == 2 else f"={L(C_CNT)}{r - 1}+{L(C_MATCH)}{r}")
        # lien par formule LIEN_HYPERTEXTE: chemin Windows tel quel (accents, &, + sans encodage d'URL)
        c = ws.cell(row=r, column=C_LIEN, value=f'=HYPERLINK({L(C_REL)}{r},"{label}")')
        c.font = Font(color="0563C1", underline="single")
        n += 1
        state["i"] = n
    return wb, ws, n

def build_sheets(wb, ws, n):
    tab = Table(displayName="Inventaire", ref=ws.dimensions)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    for i, h in enumerate(HEADERS, start=1):
        ws.column_dimensions[L(i)].width = WIDTHS.get(h, 12)
    ws.column_dimensions[L(C_INDEX)].hidden = True

    ws.column_dimensions[L(C_REL)].hidden = True

    ws.column_dimensions[L(C_CNT)].hidden = True
    blue = PatternFill("solid", fgColor="9BC2E6")
    ws.conditional_formatting.add(f"{L(C_TEXTE)}2:{L(C_TEXTE)}{n + 1}",
                                  FormulaRule(formula=[highlight_formula(f"{L(C_TEXTE)}2")], fill=blue,
                                              font=Font(bold=True)))

    # ---- Feuille de recherche par mots-clés ----
    # Ligne i = i-ème fichier trouvé (compteur cumulé de l'Inventaire, recherche dichotomique: rapide).
    # Colonnes: Titre détecté, Chemin, Ouvrir, Nom... (ordre SEARCH_COLS), colonne cachée = n° de ligne.
    rs = wb.create_sheet("Recherche", 0)
    rs["A1"] = "Mots-clés :"
    rs["A1"].font = Font(bold=True, size=14, color="1F4E79")
    rs["A1"].alignment = Alignment(horizontal="right", vertical="center")
    box = rs["B1"]
    box.fill = PatternFill("solid", fgColor="FFF2CC")
    box.font = Font(bold=True, size=14, color="1F4E79")
    box.alignment = Alignment(vertical="center")
    thick = Side(style="thick", color="2F75B5")
    box.border = Border(left=thick, right=thick, top=thick, bottom=thick)
    rs.row_dimensions[1].height = 28
    rs["C1"] = "← un ou plusieurs mots (ordre libre), tous doivent être présents"
    rs["C1"].font = Font(italic=True, color="595959")
    total = f"Inventaire!${L(C_CNT)}${n + 1}"
    rs["B2"] = f'={total}&" fichier(s) trouvé(s) sur {n}"'
    rs["B2"].font = Font(bold=True, color="2F75B5")
    head = PatternFill("solid", fgColor="1F4E79")
    for i, h in enumerate(SEARCH_COLS, start=1):
        c = rs.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    c_row = len(SEARCH_COLS) + 1                      # colonne cachée: n° de ligne dans l'Inventaire
    rs.cell(row=3, column=c_row, value="Ligne")
    cnt = f"Inventaire!${L(C_CNT)}$2:${L(C_CNT)}${n + 1}"
    rel_col = f"Inventaire!${L(C_REL)}$2:${L(C_REL)}${n + 1}"
    lien_col = f"Inventaire!${L(C_LIEN)}$2:${L(C_LIEN)}${n + 1}"     # libellé "Ouvrir" / "Ouvrir dossier"
    link = Font(color="0563C1", underline="single")
    for i in range(1, n + 1):
        r = 3 + i
        h = f"${L(c_row)}{r}"
        rs.cell(row=r, column=c_row, value=f'=IF({i}>{total},"",IFERROR(MATCH({i}-0.5,{cnt},1),0)+1)')
        for j, name in enumerate(SEARCH_COLS, start=1):
            if name == "Ouvrir":
                c = rs.cell(row=r, column=j,
                            value=f'=IF({h}="","",HYPERLINK(INDEX({rel_col},{h}),INDEX({lien_col},{h})))')
                c.font = link
                continue
            k = HEADERS.index(name) + 1
            src = f"Inventaire!${L(k)}$2:${L(k)}${n + 1}"
            rs.cell(row=r, column=j, value=f'=IF({h}="","",IF(INDEX({src},{h})="","",INDEX({src},{h})))')
    for i, h in enumerate(SEARCH_COLS, start=1):
        rs.column_dimensions[L(i)].width = WIDTHS.get(h, 12)
    rs.column_dimensions["A"].width = max(WIDTHS["Titre détecté"], 60)
    rs.column_dimensions["B"].width = 60                                  # chemin complet, à copier
    rs.column_dimensions[L(c_row)].hidden = True
    rs.conditional_formatting.add(f"A4:A{n + 3}",
                                  FormulaRule(formula=[highlight_formula("A4")], fill=blue, font=Font(bold=True)))
    rs.freeze_panes = "B4"

    # ---- Feuille Résultats: même liste que Recherche, mais en texte simple (copiable) ----
    # Une feuille de texte ne se met pas à jour seule: filtre "Trouvé = 1" déjà posé, à réappliquer (Ctrl+Alt+L)
    # après avoir changé les mots-clés. Seuls "Ouvrir" (lien) et "Trouvé" sont des formules.
    res = wb.create_sheet("Résultats", 1)
    res["A1"] = "Résultats en texte (copiable) — mots-clés :"
    res["A1"].font = Font(bold=True, size=12, color="1F4E79")
    res["A1"].alignment = Alignment(horizontal="right")
    res["B1"] = '=IF(TRIM(Recherche!$B$1)="","(tous les fichiers)",Recherche!$B$1)'
    res["B1"].font = Font(bold=True, size=12, color="2F75B5")
    res["A2"] = ("Après avoir changé les mots-clés dans Recherche : Ctrl+Alt+L "
                 "(ou Données > Réappliquer) pour mettre la liste à jour.")
    res["A2"].font = Font(italic=True, color="C00000")
    cols = SEARCH_COLS + ["Trouvé"]
    for j, name in enumerate(cols, start=1):
        res.cell(row=3, column=j, value=name)
    for r in range(2, n + 2):                        # ligne r de l'Inventaire -> ligne r + 2 ici
        for j, name in enumerate(cols, start=1):
            if name == "Ouvrir":
                c = res.cell(row=r + 2, column=j,
                             value=f"=HYPERLINK(Inventaire!${L(C_REL)}{r},Inventaire!${L(C_LIEN)}{r})")
                c.font = link
            elif name == "Trouvé":
                res.cell(row=r + 2, column=j, value=f"=Inventaire!${L(C_MATCH)}{r}")
            else:
                res.cell(row=r + 2, column=j, value=ws.cell(row=r, column=HEADERS.index(name) + 1).value)
    ref = f"A3:{L(len(cols))}{n + 3}"
    tab = Table(displayName="Resultats", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    tab.autoFilter = AutoFilter(ref=ref, filterColumn=[FilterColumn(colId=len(cols) - 1, filters=Filters(filter=["1"]))])
    res.add_table(tab)
    for j, name in enumerate(cols, start=1):
        res.column_dimensions[L(j)].width = WIDTHS.get(name, 12)
    res.column_dimensions["A"].width = max(WIDTHS["Titre détecté"], 60)
    res.column_dimensions["B"].width = 60
    res.conditional_formatting.add(f"A4:A{n + 3}",
                                   FormulaRule(formula=[highlight_formula("A4")], fill=blue, font=Font(bold=True)))
    res.freeze_panes = "A4"

# Excel n'ouvre pas un lien dont l'adresse encodée (espace -> %20, é -> %C3%A9) dépasse 218 caractères
# ("Une erreur inattendue s'est produite"). Constaté sur 07-AVANT-PROJET: <= 207 ok, >= 227 en erreur.
EXCEL_LINK_MAX = 218

def _link_len(path):
    import urllib.parse
    return len(urllib.parse.quote(str(path).replace("\\", "/"), safe="/:"))

def short_path(path):
    """Chemin court Windows 8.3 (H:\\2457-R~1\\...) pour les chemins trop longs; None si indisponible
    (hors Windows, ou noms courts désactivés sur le disque, fréquent sur les lecteurs réseau)."""
    try:
        import ctypes
        buf = ctypes.create_unicode_buffer(32768)
        n = ctypes.windll.kernel32.GetShortPathNameW("\\\\?\\" + str(path), buf, 32768)
        sp = buf.value[4:] if 0 < n < 32768 and buf.value.startswith("\\\\?\\") else None
        return sp if sp and len(sp) < len(str(path)) else None
    except Exception:
        return None

SHORTCUTS = "_liens"   # dossier (caché) des raccourcis vers les fichiers aux chemins trop longs pour Excel

def prepare_shortcuts(root):
    """Vide le dossier des raccourcis d'une analyse précédente (seulement nos fichiers .url)."""
    d = root / SHORTCUTS
    if d.is_dir():
        for f in d.glob("*.url"):
            try:
                f.unlink()
            except OSError:
                pass

def make_shortcut(p, root, num):
    """Raccourci Windows (.url) au nom court vers p: Excel ouvre le raccourci (chemin court), Windows ouvre le PDF.
    Retourne le chemin relatif du raccourci, ou None si impossible (dossier en lecture seule...)."""
    d = root / SHORTCUTS
    try:
        new = not d.is_dir()
        d.mkdir(exist_ok=True)
        if new and os.name == "nt":
            import ctypes
            ctypes.windll.kernel32.SetFileAttributesW(str(d), 2)        # dossier caché
        f = d / f"{num:05d}.url"
        f.write_text(f"[InternetShortcut]\r\nURL={p.as_uri()}\r\n", encoding="ascii")
        rel = os.path.relpath(f, root)
        return rel if _link_len(f) <= EXCEL_LINK_MAX else None
    except (OSError, ValueError, UnicodeEncodeError):
        return None

def link_for(p, root, num=0):
    """Cible du lien "Ouvrir" + libellé. Relatif au fichier Excel (enregistré dans root) si l'adresse complète tient
    dans la limite d'Excel; sinon chemin court 8.3; sinon raccourci .url au nom court; en dernier recours le dossier
    du fichier (ou un dossier parent) qui tient dans la limite."""
    if _link_len(p) <= EXCEL_LINK_MAX:
        return os.path.relpath(p, root), "Ouvrir"
    sp = short_path(p)
    if sp and _link_len(sp) <= EXCEL_LINK_MAX:
        return sp, "Ouvrir"
    sc = make_shortcut(p, root, num)
    if sc:
        return sc, "Ouvrir"
    d = p.parent
    while _link_len(d) > EXCEL_LINK_MAX and d != root and d.parent != d:
        d = d.parent
    return os.path.relpath(d, root), "Ouvrir dossier"

def absolute_links(wb, root):
    """Excel enregistré ailleurs que dans le dossier analysé (repli): les liens relatifs ne marcheraient plus."""
    ws = wb["Inventaire"]
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=C_REL)
        if c.value and not os.path.isabs(c.value):
            c.value = clean(str(root / c.value))          # la formule de la colonne Lien suit cette cellule

def default_out(root):
    """Fichier Excel enregistré dans le dossier analysé, nommé comme ce dossier (ex: D:\\Projet\\DCE -> DCE.xlsx)."""
    return root / f"{root.name or 'Inventaire'}.xlsx"

def save_workbook(wb, out, root):
    """Sauvegarde avec solutions de repli si le dossier n'existe pas, est protégé ou le fichier est verrouillé."""
    home = pathlib.Path.home()
    stamped = out.with_name(f"{out.stem}_{STAMP}{out.suffix}")    # fichier déjà ouvert dans Excel -> nom daté
    for target in [out, stamped, home / "OneDrive" / "Desktop" / out.name, home / "Desktop" / out.name,
                   home / "Documents" / out.name, home / out.name]:
        try:
            if target.parent.resolve() != root.resolve():
                absolute_links(wb, root)
            wb.save(target)
            return target
        except OSError:          # PermissionError, FileNotFoundError...
            continue
    raise PermissionError("Impossible d'enregistrer le fichier Excel (droits insuffisants).")

def fmt_time(sec):
    sec = int(sec)
    return f"{sec // 3600}h {sec % 3600 // 60:02d}min" if sec >= 3600 else f"{sec // 60}min {sec % 60:02d}s"

def run_job(root, out, state, cancel):
    """Tâche complète (thread): analyse -> feuilles -> sauvegarde."""
    try:
        pdfs = list_pdfs(root)
        state["total"] = len(pdfs)
        if not pdfs:
            state["result"] = (0, None, False)
            return
        wb, ws, n = scan(root, pdfs, state, cancel)
        state["phase"] = "save"
        if n:
            build_sheets(wb, ws, n)
            saved = save_workbook(wb, out, root)
        else:
            saved = None
        state["result"] = (n, saved, cancel.is_set())
    except Exception as e:
        state["error"] = f"{type(e).__name__}: {e}"
    finally:
        state["done"] = True

# ------------------------------------------------------------------ Lancement ----------
if len(sys.argv) > 2 and sys.argv[1] == "--debug":
    debug(sys.argv[2], int(sys.argv[3]) - 1 if len(sys.argv) > 3 else 0)
    raise SystemExit

if GUI:
    import threading, time
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    win = tk.Tk()
    win.withdraw()
    folder = filedialog.askdirectory(title="Choisir le dossier à analyser")
    if not folder:
        raise SystemExit
    ROOT = pathlib.Path(folder)
    OUT = default_out(ROOT)

    state = {"i": 0, "total": 0, "name": "", "phase": "scan", "done": False, "result": None, "error": None}
    cancel = threading.Event()
    t0 = time.time()

    win.title("Inventaire PDF")
    win.geometry("620x230")
    win.resizable(False, False)
    win.attributes("-topmost", True)
    frm = ttk.Frame(win, padding=16)
    frm.pack(fill="both", expand=True)
    ttk.Label(frm, text=f"Dossier : {ROOT}", wraplength=580).pack(anchor="w")
    lbl_count = ttk.Label(frm, text="Recherche des PDF...", font=("Segoe UI", 11, "bold"))
    lbl_count.pack(anchor="w", pady=(12, 4))
    bar = ttk.Progressbar(frm, orient="horizontal", mode="indeterminate", length=588)
    bar.pack(fill="x")
    bar.start(12)
    lbl_file = ttk.Label(frm, text="", wraplength=580)
    lbl_file.pack(anchor="w", pady=(8, 0))
    lbl_time = ttk.Label(frm, text="")
    lbl_time.pack(anchor="w", pady=(4, 0))

    def on_cancel():
        cancel.set()
        btn.config(state="disabled", text="Arrêt en cours...")
    btn = ttk.Button(frm, text="Annuler (enregistre ce qui est fait)", command=on_cancel)
    btn.pack(anchor="e", pady=(10, 0))
    win.protocol("WM_DELETE_WINDOW", on_cancel)
    win.deiconify()

    def poll():
        i, total = state["i"], state["total"]
        el = time.time() - t0
        if state["phase"] == "save":
            bar.config(mode="indeterminate")
            bar.start(12)
            lbl_count.config(text="Création du fichier Excel...")
            lbl_file.config(text="")
            btn.config(state="disabled")
        elif total:
            if str(bar.cget("mode")) != "determinate":
                bar.stop()
                bar.config(mode="determinate", maximum=total)
            bar["value"] = i
            lbl_count.config(text=f"{i} / {total} PDF analysés  ({100 * i // total}%)")
            lbl_file.config(text=state["name"])
            eta = f"   |   reste environ {fmt_time(el / i * (total - i))}" if i >= 3 else ""
            lbl_time.config(text=f"Temps écoulé : {fmt_time(el)}{eta}")
        if state["done"]:
            bar.stop()
            win.attributes("-topmost", False)
            if state["error"]:
                messagebox.showerror("Erreur", state["error"])
            else:
                n, saved, was_cancelled = state["result"]
                if not n:
                    messagebox.showinfo("Terminé", "Aucun PDF analysé." if was_cancelled else "Aucun PDF trouvé dans ce dossier.")
                else:
                    extra = "\n(Analyse interrompue: fichier partiel)" if was_cancelled else ""
                    messagebox.showinfo("Terminé", f"{n} PDF analysés en {fmt_time(time.time() - t0)}.{extra}\nFichier : {saved}")
                    try:
                        os.startfile(saved)
                    except Exception:
                        pass
            win.destroy()
            return
        win.after(150, poll)

    threading.Thread(target=run_job, args=(ROOT, OUT, state, cancel), daemon=True).start()
    poll()
    win.mainloop()
else:
    ROOT = pathlib.Path(sys.argv[1])
    OUT = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else default_out(ROOT)
    import threading
    state = {"i": 0, "total": 0, "name": "", "phase": "scan", "done": False, "result": None, "error": None}
    run_job(ROOT, OUT, state, threading.Event())
    if state["error"]:
        raise SystemExit(state["error"])
    print(f"{state['result'][0]} PDF -> {state['result'][1]}")
